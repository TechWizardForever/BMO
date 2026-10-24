"""Deterministic engineering tools for Bemo's Engineering / Deep Reasoning role.

The LLM decides *what* to calculate; this module performs bounded, auditable math,
file inspection, plotting, CAD/mesh inspection, 3D models, and technical-reference lookup.
There is deliberately no arbitrary shell or unrestricted Python execution here. 3D models are
OpenSCAD source, a language that only describes shapes; the file-reading statements are refused.
"""

from __future__ import annotations

import ast
import csv
import io
import json
import math
import os
import re
import shutil
import socket
import statistics
import subprocess
import time
import urllib.parse
from dataclasses import dataclass
from pathlib import Path
from typing import Any


# OpenSCAD statements that read other files; everything else in the language only describes shapes.
# Checked on the whole source, comments too: a "//" inside a string could otherwise hide one.
_SCAD_FILE_ACCESS = re.compile(r"\b(?:import|surface)\s*\(|\b(?:include|use)\s*<")
# The first line of every part BMO fills in: its shape and settings, so the part can be changed later even after
# BMO slept (live 2026-10-24: "make the cup the proper size" came out solid, its 2 mm walls forgotten).
_PART_HEADER = re.compile(r"^// BMO part: (\{.*\})\s*$", re.M)


class EngineeringToolError(ValueError):
    pass


def _jsonable(value: Any) -> Any:
    """Convert common numeric/library values into JSON-safe Python types."""
    try:
        import numpy as np
        if isinstance(value, np.ndarray):
            return value.tolist()
        if isinstance(value, np.generic):
            return value.item()
    except Exception:
        pass
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, complex):
        return {"real": value.real, "imag": value.imag}
    if isinstance(value, dict):
        return {str(k): _jsonable(v) for k, v in value.items()}
    if isinstance(value, (list, tuple, set)):
        return [_jsonable(v) for v in value]
    if isinstance(value, float):
        if math.isnan(value):
            return "NaN"
        if math.isinf(value):
            return "Infinity" if value > 0 else "-Infinity"
    return value


def _dumps(value: Any) -> str:
    return json.dumps(_jsonable(value), indent=2, ensure_ascii=False, default=str)


def _f(inputs: dict[str, Any], key: str, default: float | None = None) -> float:
    value = inputs.get(key, default)
    if value is None:
        raise EngineeringToolError(f"Missing required input: {key}")
    try:
        return float(value)
    except (TypeError, ValueError) as exc:
        raise EngineeringToolError(f"{key} must be numeric") from exc


def _positive(value: float, name: str, allow_zero: bool = False) -> float:
    if (value < 0 if allow_zero else value <= 0):
        op = ">= 0" if allow_zero else "> 0"
        raise EngineeringToolError(f"{name} must be {op}")
    return value


def _fraction(value: float, name: str, allow_zero: bool = False) -> float:
    low_ok = value >= 0 if allow_zero else value > 0
    if not low_ok or value > 1:
        raise EngineeringToolError(f"{name} must be {'0..1' if allow_zero else '>0 and <=1'}")
    return value


# ---------------------------------------------------------------------------
# Unit conversion
# ---------------------------------------------------------------------------
# factor converts the unit to a canonical SI-like base for its dimension.
_UNIT_TABLE: dict[str, tuple[str, float]] = {}


def _add_units(dimension: str, mapping: dict[str, float]) -> None:
    for name, factor in mapping.items():
        _UNIT_TABLE[name.lower().replace(" ", "")] = (dimension, factor)


_add_units("length", {
    "mm": 1e-3, "millimeter": 1e-3, "millimeters": 1e-3,
    "cm": 1e-2, "m": 1.0, "meter": 1.0, "meters": 1.0,
    "km": 1000.0, "in": 0.0254, "inch": 0.0254, "inches": 0.0254,
    "ft": 0.3048, "foot": 0.3048, "feet": 0.3048, "yd": 0.9144,
    "mi": 1609.344, "mile": 1609.344, "miles": 1609.344,
})
_add_units("area", {
    "mm2": 1e-6, "mm^2": 1e-6, "cm2": 1e-4, "cm^2": 1e-4,
    "m2": 1.0, "m^2": 1.0, "in2": 0.00064516, "in^2": 0.00064516,
    "ft2": 0.09290304, "ft^2": 0.09290304,
})
_add_units("volume", {
    "mm3": 1e-9, "mm^3": 1e-9, "cm3": 1e-6, "cm^3": 1e-6, "cc": 1e-6,
    "ml": 1e-6, "l": 1e-3, "liter": 1e-3, "litre": 1e-3,
    "m3": 1.0, "m^3": 1.0, "in3": 1.6387064e-5, "in^3": 1.6387064e-5,
    "ft3": 0.028316846592, "ft^3": 0.028316846592,
    "gal": 0.003785411784, "usgal": 0.003785411784,
})
_add_units("mass", {"mg": 1e-6, "g": 1e-3, "kg": 1.0, "oz": 0.028349523125, "lb": 0.45359237})
_add_units("force", {"n": 1.0, "kn": 1000.0, "lbf": 4.4482216152605})
_add_units("torque", {
    "nm": 1.0, "n*m": 1.0, "n-m": 1.0,
    "lbfft": 1.3558179483314, "lbf*ft": 1.3558179483314, "ftlb": 1.3558179483314,
    "lbfin": 0.11298482902762, "lbf*in": 0.11298482902762, "inlb": 0.11298482902762,
})
_add_units("pressure", {
    "pa": 1.0, "kpa": 1e3, "mpa": 1e6, "gpa": 1e9, "bar": 1e5,
    "psi": 6894.757293168, "ksi": 6_894_757.293168,
})
_add_units("energy", {"j": 1.0, "kj": 1e3, "mj": 1e6, "wh": 3600.0, "kwh": 3.6e6, "btu": 1055.05585262})
_add_units("power", {"w": 1.0, "kw": 1e3, "mw": 1e6, "hp": 745.6998715822702})
_add_units("speed", {
    "m/s": 1.0, "mps": 1.0, "km/h": 1/3.6, "kph": 1/3.6,
    "mph": 0.44704, "ft/s": 0.3048, "fps": 0.3048,
})
_add_units("angular_speed", {"rpm": 1.0, "rev/min": 1.0})
_add_units("angle", {"rad": 1.0, "radian": 1.0, "deg": math.pi/180.0, "degree": math.pi/180.0})
_add_units("time", {"s": 1.0, "sec": 1.0, "min": 60.0, "h": 3600.0, "hr": 3600.0, "day": 86400.0})
_add_units("voltage", {"v": 1.0, "mv": 1e-3, "kv": 1e3})
_add_units("current", {"a": 1.0, "ma": 1e-3, "ua": 1e-6})
_add_units("resistance", {"ohm": 1.0, "ohms": 1.0, "kohm": 1e3, "mohm": 1e6})
_add_units("capacitance", {"f": 1.0, "mf": 1e-3, "uf": 1e-6, "nf": 1e-9, "pf": 1e-12})
_add_units("charge_capacity", {"ah": 1.0, "mah": 1e-3})
_add_units("frequency", {"hz": 1.0, "khz": 1e3, "mhz": 1e6, "ghz": 1e9})


def convert_units(value: float, from_unit: str, to_unit: str) -> dict[str, Any]:
    f = from_unit.strip().lower().replace(" ", "")
    t = to_unit.strip().lower().replace(" ", "")
    # Bare F is ambiguous with farads. Treat it as Fahrenheit only when the
    # other side is unmistakably a temperature; F->uF remains capacitance.
    temp_units = {"c", "°c", "celsius", "°f", "fahrenheit", "k", "kelvin"}
    f_is_temp = f in temp_units or (f == "f" and t in temp_units)
    t_is_temp = t in temp_units or (t == "f" and f in temp_units)
    if f_is_temp or t_is_temp:
        if not f_is_temp or not t_is_temp:
            raise EngineeringToolError("Temperature units can only convert to other temperature units")
        if f in {"c", "°c", "celsius"}:
            kelvin = float(value) + 273.15
        elif f in {"f", "°f", "fahrenheit"}:
            kelvin = (float(value) - 32.0) * 5.0 / 9.0 + 273.15
        else:
            kelvin = float(value)
        if t in {"c", "°c", "celsius"}:
            out = kelvin - 273.15
        elif t in {"f", "°f", "fahrenheit"}:
            out = (kelvin - 273.15) * 9.0 / 5.0 + 32.0
        else:
            out = kelvin
        return {"value": out, "unit": to_unit, "dimension": "temperature"}
    try:
        dim_f, factor_f = _UNIT_TABLE[f]
        dim_t, factor_t = _UNIT_TABLE[t]
    except KeyError as exc:
        raise EngineeringToolError(f"Unsupported unit: {exc.args[0]}") from exc
    if dim_f != dim_t:
        raise EngineeringToolError(f"Incompatible dimensions: {dim_f} -> {dim_t}")
    out = float(value) * factor_f / factor_t
    return {"value": out, "unit": to_unit, "dimension": dim_f}


# ---------------------------------------------------------------------------
# Safe numerical calculator (no imports/files/network/shell)
# ---------------------------------------------------------------------------
_SAFE_FUNCS: dict[str, Any] | None = None
_ALLOWED_AST = (
    ast.Expression, ast.Constant, ast.Name, ast.Load, ast.BinOp, ast.UnaryOp,
    ast.Add, ast.Sub, ast.Mult, ast.Div, ast.FloorDiv, ast.Mod, ast.Pow,
    ast.UAdd, ast.USub, ast.List, ast.Tuple, ast.Dict, ast.Call, ast.keyword,
    ast.Subscript, ast.Slice, ast.Compare, ast.Eq, ast.NotEq, ast.Lt, ast.LtE,
    ast.Gt, ast.GtE, ast.BoolOp, ast.And, ast.Or, ast.IfExp,
)


def _safe_functions() -> dict[str, Any]:
    global _SAFE_FUNCS
    if _SAFE_FUNCS is not None:
        return _SAFE_FUNCS
    import numpy as np
    # NumPy 2.x added np.trapezoid; NumPy 1.x exposes the same operation as
    # np.trapz. Bemo deliberately pins NumPy 1.26.x for CadQuery
    # compatibility, so choose the available spelling at runtime instead of
    # eagerly touching a NumPy-2-only attribute while building this table.
    trapezoid_fn = getattr(np, "trapezoid", None)
    if trapezoid_fn is None:
        trapezoid_fn = np.trapz
    funcs = {
        "abs": abs, "round": round, "min": min, "max": max, "sum": sum,
        "sqrt": math.sqrt, "sin": math.sin, "cos": math.cos, "tan": math.tan,
        "asin": math.asin, "acos": math.acos, "atan": math.atan, "atan2": math.atan2,
        "exp": math.exp, "log": math.log, "log10": math.log10,
        "floor": math.floor, "ceil": math.ceil, "radians": math.radians, "degrees": math.degrees,
        "mean": np.mean, "median": np.median, "std": np.std, "variance": np.var,
        "array": np.array, "linspace": np.linspace, "arange": np.arange,
        "interp": np.interp, "polyfit": np.polyfit, "polyval": np.polyval,
        "dot": np.dot, "cross": np.cross, "matmul": np.matmul,
        "gradient": np.gradient, "trapezoid": trapezoid_fn, "percentile": np.percentile,
        "quantile": np.quantile, "clip": np.clip, "where": np.where,
        "solve_linear": np.linalg.solve, "det": np.linalg.det, "inv": np.linalg.inv,
        "norm": np.linalg.norm, "eigvals": np.linalg.eigvals,
        "pi": math.pi, "e": math.e,
    }
    _SAFE_FUNCS = funcs
    return funcs


def _safe_numeric_eval(expression: str, variables: dict[str, Any] | None = None) -> Any:
    if len(expression) > 4000:
        raise EngineeringToolError("Expression exceeds safety limit")
    tree = ast.parse(expression, mode="eval")
    funcs = _safe_functions()
    allowed_names = set(funcs) | set((variables or {}).keys())
    for node in ast.walk(tree):
        if not isinstance(node, _ALLOWED_AST):
            raise EngineeringToolError(f"Unsupported expression element: {type(node).__name__}")
        if isinstance(node, ast.Name) and node.id not in allowed_names:
            raise EngineeringToolError(f"Unknown name: {node.id}")
        if isinstance(node, ast.Call):
            if not isinstance(node.func, ast.Name) or node.func.id not in funcs or not callable(funcs[node.func.id]):
                raise EngineeringToolError("Only approved numeric functions may be called")
    env = {"__builtins__": {}}
    env.update(funcs)
    env.update(variables or {})
    return eval(compile(tree, "<engineering-expression>", "eval"), env, {})


def calculate_python(arguments: dict[str, Any]) -> dict[str, Any]:
    """Bounded numeric operations. The name is user-facing; execution is AST restricted."""
    operation = str(arguments.get("operation", "expression")).lower()
    if operation == "expression":
        value = _safe_numeric_eval(str(arguments.get("expression", "")), dict(arguments.get("variables") or {}))
        return {"operation": operation, "result": _jsonable(value)}
    if operation == "linear_solve":
        import numpy as np
        a = np.asarray(arguments.get("matrix"), dtype=float)
        b = np.asarray(arguments.get("vector"), dtype=float)
        if a.ndim != 2 or b.ndim not in (1, 2) or max(a.shape) > 100:
            raise EngineeringToolError("linear_solve expects a matrix/vector up to 100x100")
        return {"operation": operation, "solution": np.linalg.solve(a, b).tolist()}
    if operation == "polyfit":
        import numpy as np
        x = np.asarray(arguments.get("x"), dtype=float)
        y = np.asarray(arguments.get("y"), dtype=float)
        degree = int(arguments.get("degree", 1))
        if x.size != y.size or x.size > 10000 or degree < 0 or degree > 12:
            raise EngineeringToolError("Invalid polyfit inputs")
        coeff = np.polyfit(x, y, degree)
        return {"operation": operation, "degree": degree, "coefficients_high_to_low": coeff.tolist()}
    if operation == "interpolate":
        import numpy as np
        x = np.asarray(arguments.get("x"), dtype=float)
        y = np.asarray(arguments.get("y"), dtype=float)
        x_new = np.asarray(arguments.get("x_new"), dtype=float)
        if x.size != y.size or x.size > 100000 or x_new.size > 100000:
            raise EngineeringToolError("Invalid interpolation inputs")
        return {"operation": operation, "y_new": np.interp(x_new, x, y).tolist()}
    if operation in {"integrate", "root", "minimize"}:
        from scipy import integrate as sp_integrate
        from scipy import optimize as sp_optimize
        expr = str(arguments.get("expression", ""))
        variable = str(arguments.get("variable", "x"))
        constants = dict(arguments.get("variables") or {})
        if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", variable):
            raise EngineeringToolError("Invalid variable name")
        def f(x: float) -> float:
            values = dict(constants)
            values[variable] = float(x)
            return float(_safe_numeric_eval(expr, values))
        lo = _f(arguments, "x_min")
        hi = _f(arguments, "x_max")
        if not lo < hi:
            raise EngineeringToolError("x_min must be less than x_max")
        if operation == "integrate":
            result, error = sp_integrate.quad(f, lo, hi, limit=200)
            return {"operation": operation, "integral": result, "estimated_error": error}
        if operation == "root":
            root = sp_optimize.brentq(f, lo, hi, maxiter=200)
            return {"operation": operation, "root": root, "residual": f(root)}
        result = sp_optimize.minimize_scalar(f, bounds=(lo, hi), method="bounded", options={"maxiter": 500})
        return {"operation": operation, "x": float(result.x), "value": float(result.fun), "success": bool(result.success), "message": str(result.message)}
    raise EngineeringToolError(f"Unsupported calculate_python operation: {operation}")


# ---------------------------------------------------------------------------
# Symbolic math
# ---------------------------------------------------------------------------
def _sympy_parse(text: str, symbol_names: list[str]):
    # Keep symbolic parsing mathematical, not Python-like. Decimal points are
    # allowed only inside numbers; attribute access, strings, indexing, and
    # dunder names are rejected before SymPy's parser sees the text.
    if len(text) > 4000 or "__" in text:
        raise EngineeringToolError("Unsafe or excessive symbolic expression")
    if not re.fullmatch(r"[A-Za-z0-9_+\-*/^().,= \t]+", text):
        raise EngineeringToolError("Symbolic expression contains unsupported characters")
    without_decimals = re.sub(r"(?<=\d)\.(?=\d)", "", text)
    if "." in without_decimals:
        raise EngineeringToolError("Attribute access is not allowed in symbolic expressions")
    import sympy as sp
    from sympy.parsing.sympy_parser import parse_expr, standard_transformations, convert_xor
    names = {name: sp.Symbol(name) for name in symbol_names}
    local = {
        **names,
        "sin": sp.sin, "cos": sp.cos, "tan": sp.tan,
        "asin": sp.asin, "acos": sp.acos, "atan": sp.atan,
        "exp": sp.exp, "log": sp.log, "sqrt": sp.sqrt,
        "pi": sp.pi, "E": sp.E,
    }
    global_dict = {
        "__builtins__": {}, "Integer": sp.Integer, "Float": sp.Float,
        "Rational": sp.Rational, "Symbol": sp.Symbol,
    }
    return parse_expr(text, local_dict=local, global_dict=global_dict, transformations=standard_transformations + (convert_xor,), evaluate=True), names


def solve_math(arguments: dict[str, Any]) -> dict[str, Any]:
    import sympy as sp
    operation = str(arguments.get("operation", "solve")).lower()
    symbols = list(arguments.get("symbols") or [])
    if not symbols:
        symbols = sorted(set(re.findall(r"\b[A-Za-z][A-Za-z0-9_]*\b", " ".join([str(arguments.get("expression", ""))] + [str(x) for x in arguments.get("equations", [])]))) - {"sin", "cos", "tan", "exp", "log", "sqrt", "pi", "E"})
    if len(symbols) > 20 or any(not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", s) for s in symbols):
        raise EngineeringToolError("Invalid or excessive symbol list")
    if operation == "system":
        equations = list(arguments.get("equations") or [])
        if not equations:
            raise EngineeringToolError("system requires equations")
        parsed = []
        names = None
        for eq in equations:
            text = str(eq)
            if "=" in text:
                left, right = text.split("=", 1)
                l, names = _sympy_parse(left, symbols)
                r, _ = _sympy_parse(right, symbols)
                parsed.append(sp.Eq(l, r))
            else:
                e, names = _sympy_parse(text, symbols)
                parsed.append(e)
        sol = sp.solve(parsed, [names[s] for s in symbols], dict=True)
        return {"operation": operation, "solutions": [str(x) for x in sol]}
    expr_text = str(arguments.get("expression", ""))
    if not expr_text:
        raise EngineeringToolError("expression is required")
    if "=" in expr_text:
        left, right = expr_text.split("=", 1)
        left_expr, names = _sympy_parse(left, symbols)
        right_expr, _ = _sympy_parse(right, symbols)
        expr = sp.Eq(left_expr, right_expr)
    else:
        expr, names = _sympy_parse(expr_text, symbols)
    variable = str(arguments.get("variable") or (symbols[0] if symbols else "x"))
    if variable not in names:
        names[variable] = sp.Symbol(variable)
    var = names[variable]
    if operation == "solve":
        return {"operation": operation, "solutions": [str(x) for x in sp.solve(expr, var)]}
    if operation == "simplify":
        return {"operation": operation, "result": str(sp.simplify(expr))}
    if operation == "differentiate":
        order = max(1, min(8, int(arguments.get("order", 1))))
        return {"operation": operation, "result": str(sp.diff(expr, var, order))}
    if operation == "integrate":
        return {"operation": operation, "result": str(sp.integrate(expr, var))}
    if operation == "roots":
        return {"operation": operation, "roots": [str(x) for x in sp.solve(expr, var)]}
    raise EngineeringToolError(f"Unsupported solve_math operation: {operation}")


# ---------------------------------------------------------------------------
# Domain calculators
# ---------------------------------------------------------------------------
def mechanical_calculator(operation: str, inputs: dict[str, Any]) -> dict[str, Any]:
    op = operation.lower()
    if op in {"gear_ratio", "sprocket_ratio"}:
        driver = _positive(_f(inputs, "driver_teeth"), "driver_teeth")
        driven = _positive(_f(inputs, "driven_teeth"), "driven_teeth")
        ratio = driven / driver
        result = {"reduction_ratio_driven_over_driver": ratio, "speed_ratio_output_over_input": 1.0 / ratio}
        if inputs.get("input_rpm") is not None:
            result["output_rpm"] = _f(inputs, "input_rpm") / ratio
        if inputs.get("input_torque") is not None:
            eff = _fraction(_f(inputs, "efficiency", 1.0), "efficiency")
            result["output_torque_same_units_as_input"] = _f(inputs, "input_torque") * ratio * eff
        return result
    if op == "vehicle_speed":
        rpm = _positive(_f(inputs, "engine_rpm"), "engine_rpm", allow_zero=True)
        ratio = _positive(_f(inputs, "overall_ratio"), "overall_ratio")
        diameter_m = convert_units(_f(inputs, "tire_diameter"), str(inputs.get("diameter_unit", "in")), "m")["value"]
        wheel_rpm = rpm / ratio
        speed_mps = wheel_rpm * math.pi * diameter_m / 60.0
        return {"wheel_rpm": wheel_rpm, "speed_mps": speed_mps, "speed_mph": speed_mps / 0.44704, "speed_kph": speed_mps * 3.6}
    if op == "rpm_from_speed":
        speed_mps = convert_units(_f(inputs, "speed"), str(inputs.get("speed_unit", "mph")), "m/s")["value"]
        ratio = _positive(_f(inputs, "overall_ratio"), "overall_ratio")
        diameter_m = convert_units(_f(inputs, "tire_diameter"), str(inputs.get("diameter_unit", "in")), "m")["value"]
        wheel_rpm = speed_mps * 60.0 / (math.pi * diameter_m)
        return {"wheel_rpm": wheel_rpm, "engine_rpm": wheel_rpm * ratio}
    if op == "torque_multiplication":
        torque = _f(inputs, "input_torque")
        ratio = _positive(_f(inputs, "ratio"), "ratio")
        eff = _fraction(_f(inputs, "efficiency", 1.0), "efficiency")
        return {"output_torque_same_units_as_input": torque * ratio * eff, "ratio": ratio, "efficiency": eff}
    if op == "belt_chain_speed":
        diameter_m = convert_units(_f(inputs, "pitch_diameter"), str(inputs.get("diameter_unit", "in")), "m")["value"]
        rpm = _f(inputs, "rpm")
        speed_mps = math.pi * diameter_m * rpm / 60.0
        return {"linear_speed_mps": speed_mps, "linear_speed_ft_s": speed_mps / 0.3048}
    if op == "pulley_cvt_ratio":
        driver = _positive(_f(inputs, "driver_pitch_diameter"), "driver_pitch_diameter")
        driven = _positive(_f(inputs, "driven_pitch_diameter"), "driven_pitch_diameter")
        reduction = driven / driver
        out = {"reduction_ratio_driven_over_driver": reduction, "speed_ratio_output_over_input": 1 / reduction}
        if inputs.get("input_rpm") is not None:
            out["output_rpm"] = _f(inputs, "input_rpm") / reduction
        return out
    if op == "beam":
        mode = str(inputs.get("mode", "simply_supported_center_load"))
        length = _positive(_f(inputs, "length_m"), "length_m")
        e = _positive(_f(inputs, "elastic_modulus_pa"), "elastic_modulus_pa")
        shape = str(inputs.get("section", "rectangular"))
        if shape == "rectangular":
            b = _positive(_f(inputs, "width_m"), "width_m")
            h = _positive(_f(inputs, "height_m"), "height_m")
            i = b * h**3 / 12.0
            c = h / 2.0
        elif shape == "circular":
            d = _positive(_f(inputs, "diameter_m"), "diameter_m")
            i = math.pi * d**4 / 64.0
            c = d / 2.0
        else:
            raise EngineeringToolError("beam section must be rectangular or circular")
        if mode == "simply_supported_center_load":
            p = _f(inputs, "load_n")
            mmax = p * length / 4.0
            delta = p * length**3 / (48.0 * e * i)
        elif mode == "cantilever_end_load":
            p = _f(inputs, "load_n")
            mmax = p * length
            delta = p * length**3 / (3.0 * e * i)
        elif mode == "simply_supported_uniform_load":
            w = _f(inputs, "distributed_load_n_per_m")
            mmax = w * length**2 / 8.0
            delta = 5.0 * w * length**4 / (384.0 * e * i)
        elif mode == "cantilever_uniform_load":
            w = _f(inputs, "distributed_load_n_per_m")
            mmax = w * length**2 / 2.0
            delta = w * length**4 / (8.0 * e * i)
        else:
            raise EngineeringToolError("Unsupported beam mode")
        return {"second_moment_area_m4": i, "max_moment_nm": mmax, "max_bending_stress_pa": mmax * c / i, "max_deflection_m": delta}
    if op == "shaft_torsion":
        torque = _f(inputs, "torque_nm")
        length = _positive(_f(inputs, "length_m"), "length_m")
        g = _positive(_f(inputs, "shear_modulus_pa"), "shear_modulus_pa")
        do = _positive(_f(inputs, "outer_diameter_m"), "outer_diameter_m")
        di = _positive(_f(inputs, "inner_diameter_m", 0.0), "inner_diameter_m", allow_zero=True)
        if di >= do:
            raise EngineeringToolError("inner_diameter_m must be smaller than outer_diameter_m")
        j = math.pi * (do**4 - di**4) / 32.0
        tau = abs(torque) * (do / 2.0) / j
        angle = torque * length / (j * g)
        return {"polar_moment_m4": j, "max_shear_stress_pa": tau, "twist_rad": angle, "twist_deg": math.degrees(angle)}
    if op == "bearing_l10":
        c = _positive(_f(inputs, "dynamic_capacity_n"), "dynamic_capacity_n")
        p = _positive(_f(inputs, "equivalent_load_n"), "equivalent_load_n")
        rpm = _positive(_f(inputs, "rpm"), "rpm")
        kind = str(inputs.get("bearing_type", "ball")).lower()
        exponent = 3.0 if kind == "ball" else 10.0 / 3.0 if kind == "roller" else _f(inputs, "life_exponent")
        l10_million_rev = (c / p) ** exponent
        hours = l10_million_rev * 1_000_000.0 / (60.0 * rpm)
        return {"bearing_type": kind, "life_exponent": exponent, "L10_million_revolutions": l10_million_rev, "L10_hours": hours, "note": "Basic rating life only; application factors, lubrication, contamination, temperature, misalignment, and minimum load can dominate real life."}
    if op == "bolt_preload_torque":
        d = _positive(_f(inputs, "nominal_diameter_m"), "nominal_diameter_m")
        k = _positive(_f(inputs, "nut_factor", 0.2), "nut_factor")
        if inputs.get("preload_n") is not None:
            preload = _f(inputs, "preload_n")
        else:
            area = _positive(_f(inputs, "tensile_stress_area_m2"), "tensile_stress_area_m2")
            proof = _positive(_f(inputs, "proof_stress_pa"), "proof_stress_pa")
            fraction = _fraction(_f(inputs, "proof_fraction", 0.75), "proof_fraction")
            preload = area * proof * fraction
        return {"preload_n": preload, "estimated_torque_nm": k * preload * d, "nut_factor": k, "note": "K-factor torque/preload is approximate and highly sensitive to lubrication, coatings, thread condition, and joint geometry."}
    raise EngineeringToolError(f"Unsupported mechanical operation: {operation}")


_AWG_OHM_PER_KM_CU_20C = {
    0: 0.3224, 1: 0.4066, 2: 0.5127, 3: 0.6465, 4: 0.8152, 5: 1.028, 6: 1.296,
    7: 1.634, 8: 2.061, 9: 2.599, 10: 3.277, 11: 4.132, 12: 5.211, 13: 6.571,
    14: 8.286, 15: 10.45, 16: 13.17, 17: 16.61, 18: 20.95, 19: 26.42, 20: 33.31,
    21: 42.00, 22: 52.96, 23: 66.79, 24: 84.22, 25: 106.2, 26: 133.9, 27: 168.9,
    28: 212.9, 29: 268.5, 30: 338.6,
}


def electrical_calculator(operation: str, inputs: dict[str, Any]) -> dict[str, Any]:
    op = operation.lower()
    if op == "ohms_law":
        vals = {k: (None if inputs.get(k) is None else float(inputs[k])) for k in ("voltage_v", "current_a", "resistance_ohm", "power_w")}
        known = sum(v is not None for v in vals.values())
        if known < 2:
            raise EngineeringToolError("Provide at least two of voltage_v, current_a, resistance_ohm, power_w")
        v, i, r, p = vals["voltage_v"], vals["current_a"], vals["resistance_ohm"], vals["power_w"]
        if v is not None and i is not None:
            r = v / i if r is None and i != 0 else r; p = v * i if p is None else p
        elif v is not None and r is not None:
            i = v / r; p = v * i
        elif i is not None and r is not None:
            v = i * r; p = i * i * r
        elif p is not None and v is not None:
            i = p / v; r = v / i
        elif p is not None and i is not None:
            v = p / i; r = v / i
        elif p is not None and r is not None:
            i = math.sqrt(p / r); v = i * r
        return {"voltage_v": v, "current_a": i, "resistance_ohm": r, "power_w": p}
    if op == "resistance_network":
        values = [float(x) for x in inputs.get("resistances_ohm", [])]
        if not values or any(x <= 0 for x in values):
            raise EngineeringToolError("resistances_ohm must be positive values")
        topology = str(inputs.get("topology", "series")).lower()
        if topology == "series":
            req = sum(values)
        elif topology == "parallel":
            req = 1.0 / sum(1.0 / x for x in values)
        else:
            raise EngineeringToolError("topology must be series or parallel")
        return {"topology": topology, "equivalent_resistance_ohm": req}
    if op in {"voltage_drop", "wire_size_voltage_drop"}:
        current = _positive(_f(inputs, "current_a"), "current_a", allow_zero=True)
        length_m = _positive(convert_units(_f(inputs, "one_way_length"), str(inputs.get("length_unit", "ft")), "m")["value"], "one_way_length")
        conductors = int(inputs.get("conductors_in_path", 2))
        if conductors < 1 or conductors > 4:
            raise EngineeringToolError("conductors_in_path must be 1..4")
        temp_c = _f(inputs, "conductor_temp_c", 20.0)
        alpha = 0.00393
        temp_factor = 1 + alpha * (temp_c - 20.0)
        if op == "voltage_drop":
            gauge = int(inputs.get("awg"))
            if gauge not in _AWG_OHM_PER_KM_CU_20C:
                raise EngineeringToolError("Supported copper AWG range is 0..30")
            loop_r = _AWG_OHM_PER_KM_CU_20C[gauge] / 1000.0 * length_m * conductors * temp_factor
            drop = current * loop_r
            supply = inputs.get("system_voltage_v")
            pct = drop / float(supply) * 100.0 if supply else None
            return {"awg": gauge, "loop_resistance_ohm": loop_r, "voltage_drop_v": drop, "drop_percent": pct, "conductor_temp_c": temp_c}
        voltage = _positive(_f(inputs, "system_voltage_v"), "system_voltage_v")
        max_pct = _positive(_f(inputs, "max_drop_percent"), "max_drop_percent")
        candidates = []
        for gauge in sorted(_AWG_OHM_PER_KM_CU_20C, reverse=True):
            loop_r = _AWG_OHM_PER_KM_CU_20C[gauge] / 1000.0 * length_m * conductors * temp_factor
            pct = current * loop_r / voltage * 100.0
            if pct <= max_pct:
                candidates.append((gauge, pct, loop_r))
        if not candidates:
            return {"meets_target": False, "note": "No supported 0..30 AWG size met the voltage-drop target."}
        # Largest gauge number = thinnest acceptable conductor.
        gauge, pct, loop_r = max(candidates, key=lambda x: x[0])
        return {"meets_target": True, "minimum_size_by_voltage_drop_awg": gauge, "estimated_drop_percent": pct, "loop_resistance_ohm": loop_r, "note": "Voltage-drop sizing only. Verify ampacity, insulation temperature, bundling, ambient, duty cycle, fusing, and applicable electrical code separately."}
    if op == "battery_runtime":
        ah = _positive(_f(inputs, "capacity_ah"), "capacity_ah")
        voltage = _positive(_f(inputs, "nominal_voltage_v"), "nominal_voltage_v")
        usable = _fraction(_f(inputs, "usable_fraction", 1.0), "usable_fraction", allow_zero=True)
        efficiency = _fraction(_f(inputs, "efficiency", 1.0), "efficiency")
        wh = ah * voltage * usable
        if inputs.get("load_w") is not None:
            runtime = wh * efficiency / _positive(_f(inputs, "load_w"), "load_w")
            return {"nominal_wh": ah * voltage, "usable_wh_before_conversion_loss": wh, "runtime_hours": runtime}
        load_a = _positive(_f(inputs, "load_a"), "load_a")
        return {"nominal_wh": ah * voltage, "runtime_hours_simple_ah_basis": ah * usable * efficiency / load_a, "note": "Simple capacity model; real runtime depends on discharge rate, temperature, BMS limits, aging, and voltage cutoff."}
    if op == "solar_array":
        ns = int(inputs.get("series_modules", 1)); np_ = int(inputs.get("parallel_strings", 1))
        if ns < 1 or np_ < 1:
            raise EngineeringToolError("series_modules and parallel_strings must be >= 1")
        voc = _f(inputs, "module_voc_v"); vmp = _f(inputs, "module_vmp_v"); isc = _f(inputs, "module_isc_a"); imp = _f(inputs, "module_imp_a")
        return {"array_voc_v_stc": voc * ns, "array_vmp_v_stc": vmp * ns, "array_isc_a_stc": isc * np_, "array_imp_a_stc": imp * np_, "array_power_w_stc": vmp * imp * ns * np_, "note": "STC combination only. Controller voltage margin must use cold-temperature Voc correction from the module datasheet."}
    if op == "dc_dc_sizing":
        vin = _positive(_f(inputs, "input_voltage_v"), "input_voltage_v")
        vout = _positive(_f(inputs, "output_voltage_v"), "output_voltage_v")
        iout = _positive(_f(inputs, "output_current_a"), "output_current_a")
        eff = _fraction(_f(inputs, "efficiency", 0.9), "efficiency")
        margin = _positive(_f(inputs, "capacity_margin", 1.25), "capacity_margin")
        pout = vout * iout
        pin = pout / eff
        return {"output_power_w": pout, "estimated_input_power_w": pin, "estimated_input_current_a": pin / vin, "recommended_converter_output_current_a": iout * margin, "recommended_converter_power_w": pout * margin, "efficiency": eff, "capacity_margin": margin}
    if op == "rc_time_constant":
        r = _positive(_f(inputs, "resistance_ohm"), "resistance_ohm")
        c = _positive(_f(inputs, "capacitance_f"), "capacitance_f")
        tau = r * c
        t = inputs.get("time_s")
        result = {"time_constant_s": tau, "five_tau_s": 5 * tau}
        if t is not None:
            result["charge_fraction_at_time"] = 1 - math.exp(-float(t) / tau)
            result["discharge_fraction_remaining_at_time"] = math.exp(-float(t) / tau)
        return result
    if op == "capacitor_energy":
        c = _positive(_f(inputs, "capacitance_f"), "capacitance_f")
        v = _f(inputs, "voltage_v")
        return {"energy_j": 0.5 * c * v * v, "charge_coulomb": c * v}
    raise EngineeringToolError(f"Unsupported electrical operation: {operation}")


def thermal_calculator(operation: str, inputs: dict[str, Any]) -> dict[str, Any]:
    op = operation.lower()
    if op == "temperature":
        return convert_units(_f(inputs, "value"), str(inputs.get("from_unit", "C")), str(inputs.get("to_unit", "F")))
    if op == "heat_energy":
        mass = _positive(_f(inputs, "mass_kg"), "mass_kg")
        cp = _positive(_f(inputs, "specific_heat_j_kgk"), "specific_heat_j_kgk")
        dt = _f(inputs, "delta_t_k")
        return {"heat_j": mass * cp * dt}
    if op == "thermal_resistance":
        if inputs.get("delta_t_k") is not None and inputs.get("heat_w") is not None:
            return {"thermal_resistance_k_per_w": _f(inputs, "delta_t_k") / _positive(_f(inputs, "heat_w"), "heat_w")}
        values = [float(x) for x in inputs.get("series_resistances_k_per_w", [])]
        if not values:
            raise EngineeringToolError("Provide delta_t_k+heat_w or series_resistances_k_per_w")
        if any(x < 0 for x in values):
            raise EngineeringToolError("Thermal resistances cannot be negative")
        return {"thermal_resistance_k_per_w": sum(values)}
    if op == "steady_state_rise":
        power = _f(inputs, "power_w")
        rth = _positive(_f(inputs, "thermal_resistance_k_per_w"), "thermal_resistance_k_per_w")
        rise = power * rth
        result = {"temperature_rise_k": rise}
        if inputs.get("ambient_c") is not None:
            result["estimated_temperature_c"] = _f(inputs, "ambient_c") + rise
        return result
    if op == "conduction":
        length = _positive(_f(inputs, "thickness_m"), "thickness_m")
        k = _positive(_f(inputs, "conductivity_w_mk"), "conductivity_w_mk")
        area = _positive(_f(inputs, "area_m2"), "area_m2")
        r = length / (k * area)
        result = {"thermal_resistance_k_per_w": r}
        if inputs.get("delta_t_k") is not None:
            result["heat_flow_w"] = _f(inputs, "delta_t_k") / r
        return result
    if op == "convection":
        h = _positive(_f(inputs, "h_w_m2k"), "h_w_m2k")
        area = _positive(_f(inputs, "area_m2"), "area_m2")
        r = 1.0 / (h * area)
        result = {"thermal_resistance_k_per_w": r}
        if inputs.get("delta_t_k") is not None:
            result["heat_flow_w"] = _f(inputs, "delta_t_k") / r
        return result
    raise EngineeringToolError(f"Unsupported thermal operation: {operation}")


# ---------------------------------------------------------------------------
# File inspection / plots / CAD
# ---------------------------------------------------------------------------
@dataclass
class EngineeringToolkit:
    output_root: Path
    max_file_bytes: int = 50 * 1024 * 1024
    max_text_chars: int = 30_000
    cad_output_dir: Path | None = None
    last_part: dict | None = None
    openscad_binary: str = "openscad"
    openscad_timeout_seconds: float = 180.0

    def __post_init__(self):
        self.output_root = Path(self.output_root).expanduser().resolve()
        self.output_root.mkdir(parents=True, exist_ok=True)

    def _path(self, raw: str) -> Path:
        path = Path(raw).expanduser().resolve()
        if not path.exists() or not path.is_file():
            raise EngineeringToolError(f"File not found: {path}")
        if path.stat().st_size > self.max_file_bytes:
            raise EngineeringToolError(f"File exceeds {self.max_file_bytes} byte engineering-tool limit")
        return path

    def inspect_file(self, arguments: dict[str, Any]) -> dict[str, Any]:
        path = self._path(str(arguments.get("path", "")))
        suffix = path.suffix.lower()
        max_rows = max(1, min(200, int(arguments.get("max_rows", 30))))
        result: dict[str, Any] = {"path": str(path), "size_bytes": path.stat().st_size, "type": suffix}
        if suffix in {".txt", ".md", ".log", ".ini", ".cfg", ".yaml", ".yml", ".scad"}:
            text = path.read_text(encoding="utf-8", errors="replace")
            result.update({"text": text[:self.max_text_chars], "truncated": len(text) > self.max_text_chars})
            return result
        if suffix == ".json":
            obj = json.loads(path.read_text(encoding="utf-8", errors="replace"))
            text = json.dumps(obj, indent=2, ensure_ascii=False)
            result.update({"json_preview": text[:self.max_text_chars], "truncated": len(text) > self.max_text_chars})
            return result
        if suffix in {".csv", ".tsv"}:
            delimiter = "\t" if suffix == ".tsv" else ","
            with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as fh:
                reader = csv.reader(fh, delimiter=delimiter)
                rows = []
                for idx, row in enumerate(reader):
                    if idx >= max_rows:
                        break
                    rows.append(row[:100])
            result["rows"] = rows
            result["row_count_preview"] = len(rows)
            return result
        if suffix in {".xlsx", ".xlsm"}:
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
            sheet_name = str(arguments.get("sheet") or wb.sheetnames[0])
            if sheet_name not in wb.sheetnames:
                raise EngineeringToolError(f"Sheet not found. Available: {wb.sheetnames}")
            ws = wb[sheet_name]
            rows = []
            for idx, row in enumerate(ws.iter_rows(values_only=True)):
                if idx >= max_rows:
                    break
                rows.append([_jsonable(v) for v in row[:100]])
            result.update({"sheets": wb.sheetnames, "sheet": sheet_name, "rows": rows, "max_row": ws.max_row, "max_column": ws.max_column})
            wb.close()
            return result
        if suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            start = max(1, int(arguments.get("start_page", 1)))
            count = max(1, min(20, int(arguments.get("page_count", 5))))
            chunks = []
            for page_num in range(start - 1, min(len(reader.pages), start - 1 + count)):
                text = reader.pages[page_num].extract_text() or ""
                chunks.append(f"--- page {page_num + 1} ---\n{text}")
                if sum(len(x) for x in chunks) >= self.max_text_chars:
                    break
            text = "\n".join(chunks)
            result.update({"pages": len(reader.pages), "start_page": start, "text": text[:self.max_text_chars], "truncated": len(text) > self.max_text_chars})
            return result
        raise EngineeringToolError("Unsupported file type. Supported: txt/md/log/scad/json/csv/tsv/xlsx/xlsm/pdf")

    def plot_data(self, arguments: dict[str, Any]) -> dict[str, Any]:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        series = list(arguments.get("series") or [])
        if not series:
            raise EngineeringToolError("series is required")
        if len(series) > 12:
            raise EngineeringToolError("At most 12 series per plot")
        fig, ax = plt.subplots()
        for item in series:
            y = list(item.get("y") or [])
            x = list(item.get("x") or range(len(y)))
            if len(x) != len(y) or len(x) > 100_000:
                raise EngineeringToolError("Each plot series needs equally sized x/y arrays <= 100000")
            ax.plot(x, y, label=str(item.get("label", "")) or None)
        if arguments.get("title"):
            ax.set_title(str(arguments["title"]))
        if arguments.get("x_label"):
            ax.set_xlabel(str(arguments["x_label"]))
        if arguments.get("y_label"):
            ax.set_ylabel(str(arguments["y_label"]))
        if any(str(item.get("label", "")) for item in series):
            ax.legend()
        ax.grid(bool(arguments.get("grid", True)))
        plot_dir = self.output_root / "plots"
        plot_dir.mkdir(parents=True, exist_ok=True)
        requested = str(arguments.get("filename") or "engineering_plot.png")
        safe = re.sub(r"[^A-Za-z0-9._-]+", "_", Path(requested).name)
        if not safe.lower().endswith(".png"):
            safe += ".png"
        out = plot_dir / safe
        if out.exists():
            stem, suffix = out.stem, out.suffix
            index = 2
            while (plot_dir / f"{stem}-{index}{suffix}").exists():
                index += 1
            out = plot_dir / f"{stem}-{index}{suffix}"
        fig.tight_layout()
        fig.savefig(out, dpi=max(72, min(300, int(arguments.get("dpi", 140)))))
        plt.close(fig)
        return {"path": str(out), "series_count": len(series)}

    def inspect_cad(self, arguments: dict[str, Any]) -> dict[str, Any]:
        path = self._path(str(arguments.get("path", "")))
        mode = str(arguments.get("mode", "inspect")).lower()
        if mode == "interference":
            other = self._path(str(arguments.get("other_path", "")))
            return self._cad_interference(path, other)
        if mode == "compare":
            other = self._path(str(arguments.get("other_path", "")))
            return self._cad_compare(path, other, int(arguments.get("sample_count", 1200)))
        if mode in {"section", "overhang", "draft", "thickness", "enclosed_volume", "mass_properties"}:
            return self._cad_advanced(path, mode, arguments)
        if path.suffix.lower() in {".stl", ".obj", ".ply", ".off", ".3mf"}:
            return self._inspect_mesh(path)
        if path.suffix.lower() in {".step", ".stp"}:
            return self._inspect_step(path)
        raise EngineeringToolError("CAD inspection supports STEP/STP and mesh STL/OBJ/PLY/OFF/3MF")

    def make_cad_model(self, arguments: dict[str, Any]) -> dict[str, Any]:
        """OpenSCAD source -> .scad + .stl (never overwriting), then the STL is measured (2026-10-20)."""
        source = str(arguments.get("source") or "").strip()
        if not source:
            raise EngineeringToolError("source (the OpenSCAD design) is required")
        if len(source) > 100_000:
            raise EngineeringToolError("OpenSCAD source is too long for one model")
        if _SCAD_FILE_ACCESS.search(source):
            raise EngineeringToolError("The design may not read other files (import, include, use, surface). "
                                       "Write the whole part in source.")
        binary = shutil.which(self.openscad_binary)
        if not binary:
            raise EngineeringToolError("OpenSCAD is not installed, so no STL can be made. Say so; don't claim a file.")
        wanted_holes = arguments.get("through_holes")
        wanted_holes = None if wanted_holes in (None, "") else int(wanted_holes)
        wanted_size = arguments.get("size_mm")
        try:
            wanted_size = sorted(float(v) for v in wanted_size) if wanted_size else None
        except (TypeError, ValueError):
            raise EngineeringToolError("size_mm must be three numbers: the part's outside width, depth and height")
        if wanted_size is not None and len(wanted_size) != 3:
            raise EngineeringToolError("size_mm must be three numbers: the part's outside width, depth and height")
        if not re.search(r"\$f[nas]\s*=", source):
            source = "$fn = 64;  // smooth curves (BMO's default)\n" + source  # live 2026-10-21: a 5 mm hole had 7 sides
        folder = Path(arguments.get("_output_dir") or self.cad_output_dir or self.output_root / "models").expanduser()
        folder.mkdir(parents=True, exist_ok=True)
        stem = re.sub(r"[^A-Za-z0-9._-]+", "_", Path(str(arguments.get("name") or "model")).stem).strip("._") or "model"
        candidate, index = stem, 2
        while (folder / f"{candidate}.stl").exists() or (folder / f"{candidate}.scad").exists():
            candidate, index = f"{stem}-{index}", index + 1
        scad, stl = folder / f"{candidate}.scad", folder / f"{candidate}.stl"
        scad.write_text(source + "\n", encoding="utf-8")
        try:
            run = subprocess.run([binary, "-o", str(stl), str(scad)], cwd=str(folder), stdin=subprocess.DEVNULL,
                                 capture_output=True, text=True, timeout=self.openscad_timeout_seconds)
        except subprocess.TimeoutExpired:
            run = None
        messages = [line.strip() for line in ((run.stderr if run else "") or "").splitlines()
                    if re.match(r"(?:ERROR|WARNING|Parser error)", line.strip())][:8]
        if run is None or run.returncode != 0 or not stl.exists() or stl.stat().st_size == 0:
            stl.unlink(missing_ok=True)
            scad.unlink(missing_ok=True)
            why = (f"took longer than {self.openscad_timeout_seconds:.0f} s" if run is None
                   else "; ".join(messages) or f"OpenSCAD exit code {run.returncode}")
            raise EngineeringToolError(f"No STL was made: {why}. Fix the source and try again.")
        measured = self._inspect_mesh(stl)
        volume = measured.get("volume_native_units_cubed")
        pieces, closed, holes = self._solid_pieces(stl)
        extents = [float(v) for v in measured["extents_native_units"]]

        def refuse(why):
            stl.unlink(missing_ok=True)
            scad.unlink(missing_ok=True)
            raise EngineeringToolError(f"No STL kept: {why} Fix the source and try again.")

        # Benchmark 2026-10-22: Engineering's L-bracket put the upright beside the base, touching only along an
        # edge: 43 x 20 x 43 mm instead of 40 x 20 x 40, two plates rather than one part, and it called it done.
        if not closed or pieces != 1:
            refuse(f"it is {pieces} separate piece(s)" + ("" if closed else ", or pieces that only touch at an edge")
                   + ", not one solid part. Pieces that belong together must overlap (by 0.01 mm or more),"
                     " not just touch.")
        if wanted_size is not None and any(abs(a - b) > 0.5 for a, b in zip(sorted(extents), wanted_size)):
            refuse(f"it measures {' x '.join(f'{v:g}' for v in (round(e, 2) for e in extents))} mm, but size_mm "
                   f"{' x '.join(f'{v:g}' for v in arguments['size_mm'])} was asked. Check each translate() and size.")
        if wanted_holes is not None and holes != wanted_holes:
            # Live 2026-10-21: "the same cube but with an 8 mm hole" came out as cube() + cylinder() with no
            # difference(): a solid cube of the right size, which the size check alone passed.
            # Benchmark 2026-10-22: holes added with no difference() (Coding), and holes that started halfway up a
            # 3 mm plate, translate([x, y, 1.5]) (Engineering); the second needed the start below the part spelt out.
            hint = (" A hole must be cut out: put the part first inside difference() and the hole shapes after it."
                    " Each hole shape must start below the part and end above it, e.g. for a 3 mm plate"
                    " translate([x, y, -1]) cylinder(h=5, d=4)." if holes < wanted_holes else "")
            refuse(f"it has {holes} hole(s) going all the way through, but through_holes={wanted_holes} was "
                   f"asked.{hint}")
        box = extents[0] * extents[1] * extents[2]
        return {
            "stl_path": str(stl), "scad_path": str(scad), "units": "mm",
            "size_mm": [round(float(v), 3) for v in measured["extents_native_units"]],
            "bounds_mm": measured["bounds_native_units"],
            "volume_mm3": round(float(volume), 3) if volume is not None else None,
            "through_holes": holes,
            "solid_fraction": round(float(volume) / box, 3) if volume is not None and box > 0 else None,
            "watertight": measured["watertight"], "openscad_warnings": messages,
            "check": "Compare size_mm, through_holes and volume_mm3 with what was asked (solid_fraction 1.0 means a "
                     "plain solid block); if anything is wrong, fix the source and make it again. Tell the owner "
                     "the measured size and where the STL is saved.",
        }

    def make_cad_part(self, arguments: dict[str, Any]) -> dict[str, Any]:
        """One basic shape filled in by number (systems/cad_shapes.py); BMO writes the OpenSCAD and the size, hole
        and one-piece checks of make_cad_model apply."""
        from systems import cad_shapes
        values = dict(arguments.get("settings") or {}) if isinstance(arguments.get("settings"), dict) else {}
        # The numbers may also come beside shape and name instead of inside settings.
        values.update({k: v for k, v in arguments.items() if k not in {"shape", "kind", "name", "settings", "_output_dir"}})
        shape = arguments.get("shape") or arguments.get("kind")
        try:
            source, size, holes = cad_shapes.build(shape, values)
        except cad_shapes.ShapeError as exc:
            raise EngineeringToolError(str(exc)) from None
        kind = cad_shapes.kind_for(shape)
        # "<" and "(" escaped so a text part's words can never look like a file read to make_cad_model.
        header = "// BMO part: " + json.dumps({"shape": kind, "settings": values}, sort_keys=True, default=str
                                              ).replace("<", "\\u003c").replace("(", "\\u0028")
        try:
            made = self.make_cad_model({"source": header + "\n$fn = 64;\n" + source, "name": arguments.get("name") or shape,
                                        "size_mm": size, "through_holes": holes,
                                        "_output_dir": arguments.get("_output_dir")})
        except EngineeringToolError as exc:
            if cad_shapes.kind_for(shape) == "text" and ("measures" in str(exc) or "separate piece" in str(exc)):
                raise EngineeringToolError("No STL kept: the text doesn't fit on its plate. Make size smaller or "
                                           "the plate wider.") from None
            raise
        made["shape"] = cad_shapes.kind_for(shape)
        made["settings"] = values
        # Remembered for simple changes ("make the hole bigger"): the next request starts from these numbers.
        self.last_part = {"shape": made["shape"], "settings": values, "name": str(arguments.get("name") or shape),
                          "stl_path": made["stl_path"], "at": time.monotonic()}
        return made

    def saved_parts(self, folder: str | Path | None = None, limit: int = 6) -> list[dict[str, Any]]:
        """The newest parts saved in a folder: each one's shape and settings from its first line, or for a design
        without that line (older parts, or ones Engineering wrote itself) the design's own text."""
        folder = Path(folder or self.cad_output_dir or self.output_root / "models").expanduser()
        try:
            designs = sorted(folder.glob("*.scad"), key=lambda p: p.stat().st_mtime, reverse=True)[:limit]
        except OSError:
            return []
        parts = []
        for scad in designs:
            try:
                text = scad.read_text(encoding="utf-8", errors="replace")[:20000]
                age = time.time() - scad.stat().st_mtime
            except OSError:
                continue
            part = {"file": scad.with_suffix(".stl").name, "name": re.sub(r"-\d+$", "", scad.stem),
                    "stl_path": str(scad.with_suffix(".stl")), "age_seconds": age}
            found = _PART_HEADER.search(text)
            try:
                data = json.loads(found.group(1)) if found else None
            except ValueError:
                data = None
            if isinstance(data, dict) and data.get("shape"):
                part.update(shape=str(data["shape"]), settings=dict(data.get("settings") or {}))
            else:
                part["source"] = re.sub(r"\s+", " ", text).strip()[:400]
            parts.append(part)
        return parts

    def latest_part(self, folder: str | Path | None = None, within: float = 3600.0) -> dict[str, Any] | None:
        """The newest part with its settings, made in the last `within` seconds, in last_part's form."""
        for part in self.saved_parts(folder, limit=1):
            if part.get("shape") and part["age_seconds"] < within:
                return {"shape": part["shape"], "settings": part["settings"], "name": part["name"],
                        "stl_path": part["stl_path"], "at": time.monotonic() - part["age_seconds"]}
        return None

    @staticmethod
    def _solid_pieces(path: Path) -> tuple[int, bool, int | None]:
        """(separate pieces, every piece closed, holes going all the way through, i.e. the genus; None if open)."""
        import trimesh
        mesh = trimesh.load(path, force="mesh")
        bodies = mesh.split(only_watertight=False) if len(mesh.faces) else []
        closed = bool(bodies) and mesh.is_watertight and all(b.is_watertight for b in bodies)
        holes = int(sum(max(0, (2 - b.euler_number) // 2) for b in bodies)) if closed else None
        return len(bodies), closed, holes

    @classmethod
    def _through_holes(cls, path: Path) -> int | None:
        return cls._solid_pieces(path)[2]

    def _load_mesh_any(self, path: Path, tolerance: float = 0.15):
        """Return a trimesh representation for analysis-only operations.

        STEP is tessellated locally, so exact STEP bounds/volume/interference still
        use OCP while section/overhang/thickness can share one mesh pipeline.
        """
        import numpy as np
        import trimesh
        if path.suffix.lower() in {".stl", ".obj", ".ply", ".off", ".3mf"}:
            loaded = trimesh.load(path, force="scene")
            if isinstance(loaded, trimesh.Scene):
                geoms = [g for g in loaded.geometry.values() if hasattr(g, "vertices")]
                if not geoms:
                    raise EngineeringToolError("No mesh geometry found")
                return trimesh.util.concatenate(tuple(geoms)), "mesh-native"
            return loaded, "mesh-native"
        if path.suffix.lower() not in {".step", ".stp"}:
            raise EngineeringToolError("Unsupported CAD format")
        import cadquery as cq
        wp = cq.importers.importStep(str(path))
        vertices = []
        faces = []
        for shape in [x for x in wp.vals() if hasattr(x, "tessellate")]:
            try:
                verts, tris = shape.tessellate(max(0.01, min(2.0, float(tolerance))))
            except Exception:
                continue
            base = len(vertices)
            vertices.extend([[float(v.x), float(v.y), float(v.z)] for v in verts])
            faces.extend([[base + int(a), base + int(b), base + int(c)] for a, b, c in tris])
        if not vertices or not faces:
            raise EngineeringToolError("STEP tessellation produced no mesh")
        mesh = trimesh.Trimesh(vertices=np.asarray(vertices), faces=np.asarray(faces), process=True)
        try:
            mesh.merge_vertices()
            mesh.remove_unreferenced_vertices()
        except Exception:
            pass
        return mesh, "STEP-tessellation-native-units"

    @staticmethod
    def _unit_scale_to_m(unit: str) -> float:
        table = {"mm": 1e-3, "cm": 1e-2, "m": 1.0, "in": 0.0254, "ft": 0.3048}
        key = str(unit).lower().strip()
        if key not in table:
            raise EngineeringToolError("native_length_unit must be one of mm/cm/m/in/ft")
        return table[key]

    def _cad_advanced(self, path: Path, mode: str, arguments: dict[str, Any]) -> dict[str, Any]:
        import numpy as np
        import trimesh
        mesh, backend = self._load_mesh_any(path, float(arguments.get("tessellation_tolerance", 0.15)))
        if mode == "section":
            origin = np.asarray(arguments.get("plane_origin", [0, 0, 0]), dtype=float)
            normal = np.asarray(arguments.get("plane_normal", [0, 0, 1]), dtype=float)
            if origin.shape != (3,) or normal.shape != (3,) or np.linalg.norm(normal) == 0:
                raise EngineeringToolError("plane_origin and plane_normal must be 3-element vectors")
            normal = normal / np.linalg.norm(normal)
            section = mesh.section(plane_origin=origin, plane_normal=normal)
            if section is None:
                return {"path": str(path), "mode": mode, "intersects": False, "plane_origin": origin.tolist(), "plane_normal": normal.tolist(), "backend": backend}
            try:
                planar, _ = section.to_planar()
                bounds = planar.bounds.tolist()
                length = float(planar.length)
                entities = len(planar.entities)
            except Exception:
                bounds = None; length = float(section.length); entities = len(section.entities)
            return {"path": str(path), "mode": mode, "intersects": True, "plane_origin": origin.tolist(), "plane_normal": normal.tolist(), "section_total_length_native_units": length, "section_2d_bounds_native_units": bounds, "section_entities": entities, "backend": backend, "note": "Section geometry is derived from tessellated geometry for STEP and is suitable for inspection, not metrology-grade manufacturing release."}
        if mode == "overhang":
            build = np.asarray(arguments.get("build_direction", [0, 0, 1]), dtype=float)
            if build.shape != (3,) or np.linalg.norm(build) == 0:
                raise EngineeringToolError("build_direction must be a nonzero 3-vector")
            build = build / np.linalg.norm(build)
            threshold = max(0.0, min(89.9, float(arguments.get("max_overhang_from_vertical_deg", 45.0))))
            # Down-facing faces whose normals are more horizontal than the
            # allowed threshold are a conservative support-risk screen.
            normals = mesh.face_normals
            areas = mesh.area_faces
            dot = normals @ build
            downward = dot < 0
            angle_from_down = np.degrees(np.arccos(np.clip(-dot, -1.0, 1.0)))
            risky = downward & (angle_from_down < (90.0 - threshold))
            risky_area = float(areas[risky].sum()) if len(areas) else 0.0
            total_area = float(areas.sum()) if len(areas) else 0.0
            return {"path": str(path), "mode": mode, "build_direction": build.tolist(), "threshold_deg": threshold, "risky_face_count": int(risky.sum()), "face_count": int(len(normals)), "risky_area_native_units_squared": risky_area, "risky_area_percent": (100.0 * risky_area / total_area if total_area else 0.0), "backend": backend, "warning": "This is a geometric FDM support-risk screen. Bridges, local cooling, nozzle width, and slicer behavior are not modeled."}
        if mode == "draft":
            pull = np.asarray(arguments.get("pull_direction", [0, 0, 1]), dtype=float)
            if pull.shape != (3,) or np.linalg.norm(pull) == 0:
                raise EngineeringToolError("pull_direction must be a nonzero 3-vector")
            pull = pull / np.linalg.norm(pull)
            minimum = max(0.0, min(45.0, float(arguments.get("minimum_draft_deg", 1.0))))
            normals = np.asarray(mesh.face_normals, dtype=float)
            areas = np.asarray(mesh.area_faces, dtype=float)
            dot = np.clip(normals @ pull, -1.0, 1.0)
            # Draft of a side wall is the tilt of its normal out of the plane
            # perpendicular to pull. Caps near +/- pull are excluded because
            # they are not meaningful draft-wall candidates.
            signed = np.degrees(np.arcsin(dot))
            side = np.abs(dot) < math.cos(math.radians(15.0))
            low = side & (np.abs(signed) < minimum)
            side_area = float(areas[side].sum()) if len(areas) else 0.0
            low_area = float(areas[low].sum()) if len(areas) else 0.0
            values = signed[side]
            return {
                "path": str(path), "mode": mode, "pull_direction": pull.tolist(),
                "minimum_requested_draft_deg": minimum,
                "side_face_count": int(side.sum()), "low_draft_face_count": int(low.sum()),
                "side_area_native_units_squared": side_area, "low_draft_area_native_units_squared": low_area,
                "low_draft_area_percent_of_side_area": (100.0 * low_area / side_area if side_area else 0.0),
                "draft_signed_deg_min": (float(values.min()) if len(values) else None),
                "draft_signed_deg_max": (float(values.max()) if len(values) else None),
                "backend": backend,
                "warning": "This is a tessellated geometric draft screen, not a moldability guarantee. Parting direction, shutoffs, lifters, texture allowance, and true undercuts require design context."
            }
        if mode == "enclosed_volume":
            components = list(mesh.split(only_watertight=False))
            rows = []
            total = 0.0
            for idx, component in enumerate(components[:200]):
                watertight = bool(component.is_watertight)
                volume = float(abs(component.volume)) if watertight else None
                if volume is not None:
                    total += volume
                rows.append({
                    "component_index": idx, "watertight": watertight,
                    "volume_native_units_cubed": volume,
                    "bounds_native_units": component.bounds.tolist(),
                    "faces": int(len(component.faces)),
                })
            return {
                "path": str(path), "mode": mode, "component_count": len(components),
                "watertight_component_count": sum(1 for x in rows if x["watertight"]),
                "total_enclosed_volume_native_units_cubed": total, "components": rows,
                "backend": backend,
                "warning": "Reports closed connected mesh shells as enclosed volumes. It does not distinguish intended cavities from separate solid bodies without assembly/design context."
            }
        if mode == "thickness":
            if not mesh.is_watertight:
                raise EngineeringToolError("Wall-thickness sampling requires a watertight body")
            requested = max(20, min(1500, int(arguments.get("sample_count", 300))))
            # Keep the brute-force, dependency-free ray screen bounded. The
            # implementation deliberately avoids requiring libspatialindex/rtree
            # on small offline Bemo machines.
            face_count = max(1, len(mesh.faces))
            count = min(requested, max(20, int(20_000_000 / face_count)))
            points, face_index = trimesh.sample.sample_surface(mesh, count)
            triangles = np.asarray(mesh.triangles, dtype=float)
            normals = np.asarray(mesh.face_normals[face_index], dtype=float)
            eps = max(1e-9, float(np.linalg.norm(mesh.extents)) * 1e-9)

            def ray_distance(origin, direction):
                # Vectorized Moller-Trumbore against every triangle. Ray points
                # inward from the sampled exterior face; the first non-self hit
                # approximates local through-wall thickness.
                v0 = triangles[:, 0]; e1 = triangles[:, 1] - v0; e2 = triangles[:, 2] - v0
                h = np.cross(np.broadcast_to(direction, e2.shape), e2)
                a = np.einsum("ij,ij->i", e1, h)
                mask = np.abs(a) > 1e-12
                f = np.zeros_like(a); f[mask] = 1.0 / a[mask]
                ss = np.broadcast_to(origin, v0.shape) - v0
                u = f * np.einsum("ij,ij->i", ss, h)
                mask &= (u >= 0.0) & (u <= 1.0)
                q = np.cross(ss, e1)
                v = f * np.einsum("ij,ij->i", np.broadcast_to(direction, q.shape), q)
                mask &= (v >= 0.0) & ((u + v) <= 1.0)
                t = f * np.einsum("ij,ij->i", e2, q)
                valid = t[mask & (t > eps)]
                return float(valid.min()) if len(valid) else float("nan")

            values = np.asarray([ray_distance(p - n * eps * 10.0, -n) for p, n in zip(points, normals)], dtype=float)
            values = values[np.isfinite(values) & (values > 0)]
            if not len(values):
                raise EngineeringToolError("No valid wall-thickness samples")
            return {"path": str(path), "mode": mode, "sample_count": int(len(values)), "requested_sample_count": requested, "minimum_native_units": float(values.min()), "p05_native_units": float(np.percentile(values, 5)), "median_native_units": float(np.median(values)), "p95_native_units": float(np.percentile(values, 95)), "maximum_native_units": float(values.max()), "backend": backend + "/bounded-bruteforce-ray", "warning": "Sampled ray thickness is approximate and can miss small local minima or complex concave paths; verify critical walls with exact CAD dimensions."}
        if mode == "mass_properties":
            if not mesh.is_watertight:
                raise EngineeringToolError("Mass properties require a watertight body")
            density = float(arguments.get("density_kg_m3", 0.0))
            if density <= 0:
                raise EngineeringToolError("density_kg_m3 must be > 0")
            scale = self._unit_scale_to_m(str(arguments.get("native_length_unit", "mm")))
            volume_m3 = abs(float(mesh.volume)) * scale ** 3
            mass = density * volume_m3
            inertia_native = np.asarray(mesh.moment_inertia, dtype=float)
            inertia_kg_m2 = inertia_native * density * scale ** 5
            return {"path": str(path), "mode": mode, "density_kg_m3": density, "native_length_unit": arguments.get("native_length_unit", "mm"), "volume_m3": volume_m3, "mass_kg": mass, "center_mass_native_units": mesh.center_mass.tolist(), "inertia_about_center_kg_m2": inertia_kg_m2.tolist(), "backend": backend, "warning": "Assumes uniform density and the stated native length unit."}
        raise EngineeringToolError(f"Unsupported advanced CAD mode: {mode}")

    def _cad_compare(self, a: Path, b: Path, sample_count: int) -> dict[str, Any]:
        import numpy as np
        import trimesh
        ma, backend_a = self._load_mesh_any(a)
        mb, backend_b = self._load_mesh_any(b)
        count = max(100, min(4000, int(sample_count)))
        pa, _ = trimesh.sample.sample_surface(ma, count)
        pb, _ = trimesh.sample.sample_surface(mb, count)
        # KD-tree symmetric sampled distance avoids requiring an external exact
        # proximity backend and gives a useful revision-change magnitude.
        from scipy.spatial import cKDTree
        da, _ = cKDTree(pb).query(pa, k=1)
        db, _ = cKDTree(pa).query(pb, k=1)
        ia = self._inspect_step(a) if a.suffix.lower() in {".step", ".stp"} else self._inspect_mesh(a)
        ib = self._inspect_step(b) if b.suffix.lower() in {".step", ".stp"} else self._inspect_mesh(b)
        va = ia.get("volume_native_units_cubed"); vb = ib.get("volume_native_units_cubed")
        return {"path_a": str(a), "path_b": str(b), "mode": "compare", "backend_a": backend_a, "backend_b": backend_b, "extents_a": ia.get("extents_native_units"), "extents_b": ib.get("extents_native_units"), "extent_delta_native_units": [float(y-x) for x,y in zip(ia.get("extents_native_units", [0,0,0]), ib.get("extents_native_units", [0,0,0]))], "volume_a_native_units_cubed": va, "volume_b_native_units_cubed": vb, "volume_delta_native_units_cubed": (float(vb-va) if va is not None and vb is not None else None), "sampled_surface_distance_mean_native_units": float(np.mean(np.concatenate([da, db]))), "sampled_surface_distance_p95_native_units": float(np.percentile(np.concatenate([da, db]),95)), "sampled_surface_distance_max_native_units": float(max(da.max(), db.max())), "sample_count_each": count, "warning": "Surface distances are sampled approximations; use exact dimensions/interference for release-critical decisions."}

    def _inspect_mesh(self, path: Path) -> dict[str, Any]:
        import trimesh
        loaded = trimesh.load(path, force="scene")
        if isinstance(loaded, trimesh.Scene):
            geometries = [g for g in loaded.geometry.values() if hasattr(g, "vertices")]
            if not geometries:
                raise EngineeringToolError("No mesh geometry found")
            mesh = trimesh.util.concatenate(tuple(geometries))
            components = len(geometries)
        else:
            mesh = loaded
            components = 1
        bounds = mesh.bounds.tolist()
        result = {
            "path": str(path), "backend": "trimesh", "components": components,
            "vertices": int(len(mesh.vertices)), "faces": int(len(mesh.faces)),
            "bounds_native_units": bounds, "extents_native_units": mesh.extents.tolist(),
            "surface_area_native_units_squared": float(mesh.area),
            "watertight": bool(mesh.is_watertight),
        }
        if mesh.is_watertight:
            result["volume_native_units_cubed"] = float(abs(mesh.volume))
            result["center_mass_native_units"] = mesh.center_mass.tolist()
        else:
            result["volume_native_units_cubed"] = None
            result["note"] = "Mesh is not watertight; volume/center of mass are not reported as authoritative."
        return result

    def _inspect_step(self, path: Path) -> dict[str, Any]:
        import cadquery as cq
        wp = cq.importers.importStep(str(path))
        shapes = [s for s in wp.vals() if hasattr(s, "BoundingBox")]
        if not shapes:
            raise EngineeringToolError("No STEP shapes found")
        mins = [float("inf")] * 3; maxs = [float("-inf")] * 3
        total_volume = 0.0; total_area = 0.0
        centers = []
        weighted_center = [0.0, 0.0, 0.0]
        weighted_volume = 0.0
        for shape in shapes:
            bb = shape.BoundingBox()
            mins[0] = min(mins[0], bb.xmin); mins[1] = min(mins[1], bb.ymin); mins[2] = min(mins[2], bb.zmin)
            maxs[0] = max(maxs[0], bb.xmax); maxs[1] = max(maxs[1], bb.ymax); maxs[2] = max(maxs[2], bb.zmax)
            try:
                volume = float(shape.Volume())
                total_volume += volume
            except Exception:
                volume = 0.0
            try: total_area += float(shape.Area())
            except Exception: pass
            try:
                c = shape.Center(); center = [float(c.x), float(c.y), float(c.z)]
                centers.append(center)
                if volume > 0:
                    for i in range(3):
                        weighted_center[i] += center[i] * volume
                    weighted_volume += volume
            except Exception: pass
        ext = [maxs[i] - mins[i] for i in range(3)]
        center_of_volume = [v / weighted_volume for v in weighted_center] if weighted_volume > 0 else None
        return {"path": str(path), "backend": "cadquery/OCP", "shape_count": len(shapes), "bounds_native_units": [mins, maxs], "extents_native_units": ext, "surface_area_native_units_squared": total_area, "volume_native_units_cubed": total_volume, "center_of_volume_native_units": center_of_volume, "shape_centers_native_units": centers[:100], "note": "STEP files commonly use millimeters, but this tool reports the file/backend native units; verify document units before treating them as mm. Center of volume assumes uniform density."}

    @staticmethod
    def _bbox_overlap(a: list[list[float]], b: list[list[float]]) -> tuple[list[float], float, list[float], float]:
        overlap = [max(0.0, min(a[1][i], b[1][i]) - max(a[0][i], b[0][i])) for i in range(3)]
        gaps = [max(0.0, max(a[0][i], b[0][i]) - min(a[1][i], b[1][i])) for i in range(3)]
        separation = math.sqrt(sum(g * g for g in gaps))
        return overlap, overlap[0] * overlap[1] * overlap[2], gaps, separation

    def _cad_interference(self, a: Path, b: Path) -> dict[str, Any]:
        ia = self._inspect_step(a) if a.suffix.lower() in {".step", ".stp"} else self._inspect_mesh(a)
        ib = self._inspect_step(b) if b.suffix.lower() in {".step", ".stp"} else self._inspect_mesh(b)
        overlap, bbox_vol, bbox_gaps, bbox_sep = self._bbox_overlap(ia["bounds_native_units"], ib["bounds_native_units"])
        result = {"path_a": str(a), "path_b": str(b), "bbox_overlap_extents_native_units": overlap, "bbox_overlap_volume_native_units_cubed": bbox_vol, "bounding_boxes_overlap": bbox_vol > 0, "bbox_axis_gaps_native_units": bbox_gaps, "bbox_minimum_separation_native_units": bbox_sep}
        if a.suffix.lower() in {".step", ".stp"} and b.suffix.lower() in {".step", ".stp"}:
            try:
                import cadquery as cq
                sa = cq.importers.importStep(str(a)).combine().val()
                sb = cq.importers.importStep(str(b)).combine().val()
                result["exact_minimum_distance_native_units"] = float(sa.distance(sb))
                if bbox_vol > 0:
                    common = sa.intersect(sb)
                    vol = float(common.Volume()) if common else 0.0
                    result["exact_intersection_volume_native_units_cubed"] = vol
                    result["exact_interference"] = vol > 1e-9
                else:
                    result["exact_intersection_volume_native_units_cubed"] = 0.0
                    result["exact_interference"] = False
                # When STEP contains multiple top-level solids, report which
                # shape pairs actually overlap. This is more actionable than a
                # single combined boolean while remaining bounded.
                shapes_a = [x for x in cq.importers.importStep(str(a)).vals() if hasattr(x, "BoundingBox")]
                shapes_b = [x for x in cq.importers.importStep(str(b)).vals() if hasattr(x, "BoundingBox")]
                pairs = []
                checked = 0
                for ia, xa in enumerate(shapes_a):
                    bba = xa.BoundingBox(); ba = [[bba.xmin,bba.ymin,bba.zmin],[bba.xmax,bba.ymax,bba.zmax]]
                    for ib, xb in enumerate(shapes_b):
                        if checked >= 200:
                            break
                        checked += 1
                        bbb = xb.BoundingBox(); bb = [[bbb.xmin,bbb.ymin,bbb.zmin],[bbb.xmax,bbb.ymax,bbb.zmax]]
                        _, pair_bbox_vol, _, _ = self._bbox_overlap(ba, bb)
                        if pair_bbox_vol <= 0:
                            continue
                        try:
                            common_pair = xa.intersect(xb)
                            pair_vol = float(common_pair.Volume()) if common_pair else 0.0
                        except Exception:
                            pair_vol = 0.0
                        if pair_vol > 1e-9:
                            pairs.append({
                                "shape_a_index": ia, "shape_b_index": ib,
                                "intersection_volume_native_units_cubed": pair_vol,
                                "shape_a_bounds_native_units": ba,
                                "shape_b_bounds_native_units": bb,
                            })
                    if checked >= 200:
                        break
                result["interfering_shape_pairs"] = pairs
                result["shape_pairs_checked"] = checked
                result["shape_pair_check_truncated"] = (len(shapes_a) * len(shapes_b) > checked)
            except Exception as exc:
                result["exact_interference"] = None
                result["exact_check_error"] = f"{type(exc).__name__}: {exc}"
        else:
            result["exact_interference"] = None
            result["note"] = "For meshes this is a bounding-box screen only; exact boolean interference requires a robust mesh boolean engine."
        return result

    # -----------------------------------------------------------------------
    # Technical references: bounded public web search/fetch with URL citations
    # -----------------------------------------------------------------------
    @staticmethod
    def _safe_public_url(url: str) -> str:
        parsed = urllib.parse.urlparse(url)
        if parsed.scheme not in {"http", "https"} or not parsed.hostname:
            raise EngineeringToolError("Only public http/https URLs are allowed")
        host = parsed.hostname.lower()
        if host in {"localhost", "localhost.localdomain"} or host.endswith(".local"):
            raise EngineeringToolError("Local/private URLs are blocked")
        try:
            infos = socket.getaddrinfo(host, parsed.port or (443 if parsed.scheme == "https" else 80), type=socket.SOCK_STREAM)
            import ipaddress
            for info in infos:
                ip = ipaddress.ip_address(info[4][0])
                if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved or ip.is_multicast:
                    raise EngineeringToolError("Local/private network destinations are blocked")
        except socket.gaierror as exc:
            raise EngineeringToolError(f"Could not resolve host: {host}") from exc
        return url

    def lookup_reference(self, arguments: dict[str, Any]) -> dict[str, Any]:
        import requests
        from bs4 import BeautifulSoup
        mode = str(arguments.get("mode", "search")).lower()
        if mode == "search":
            query = str(arguments.get("query", "")).strip()
            if not query:
                raise EngineeringToolError("query is required")
            domains = [str(x).strip() for x in arguments.get("domains", []) if str(x).strip()]
            if domains:
                query = query + " " + " OR ".join(f"site:{d}" for d in domains[:6])
            resp = requests.get(
                "https://html.duckduckgo.com/html/",
                params={"q": query},
                headers={"User-Agent": "Mozilla/5.0 (compatible; BemoEngineering/1.0)"},
                timeout=10,
            )
            resp.raise_for_status()
            soup = BeautifulSoup(resp.text, "html.parser")
            results = []
            for result in soup.select(".result"):
                a = result.select_one("a.result__a")
                if not a:
                    continue
                href = str(a.get("href") or "")
                # DDG HTML can wrap target URLs in uddg.
                parsed = urllib.parse.urlparse(href)
                target = urllib.parse.parse_qs(parsed.query).get("uddg", [href])[0]
                snippet = result.select_one(".result__snippet")
                results.append({"title": a.get_text(" ", strip=True), "url": target, "snippet": snippet.get_text(" ", strip=True) if snippet else ""})
                if len(results) >= max(1, min(8, int(arguments.get("max_results", 5)))):
                    break
            return {"query": query, "results": results, "citation_rule": "Cite the returned URL for any specification or fact used."}
        if mode == "fetch":
            url = self._safe_public_url(str(arguments.get("url", "")))
            resp = requests.get(url, headers={"User-Agent": "Mozilla/5.0 (compatible; BemoEngineering/1.0)"}, timeout=15, stream=True)
            resp.raise_for_status()
            content = resp.raw.read(min(self.max_file_bytes, 8 * 1024 * 1024) + 1)
            if len(content) > min(self.max_file_bytes, 8 * 1024 * 1024):
                raise EngineeringToolError("Reference exceeds fetch safety limit")
            ctype = (resp.headers.get("content-type") or "").lower()
            if "pdf" in ctype or url.lower().endswith(".pdf"):
                from pypdf import PdfReader
                reader = PdfReader(io.BytesIO(content))
                chunks = []
                for idx, page in enumerate(reader.pages[:20]):
                    chunks.append(f"--- page {idx + 1} ---\n{page.extract_text() or ''}")
                    if sum(len(x) for x in chunks) >= self.max_text_chars:
                        break
                text = "\n".join(chunks)
            else:
                soup = BeautifulSoup(content, "html.parser")
                for tag in soup(["script", "style", "noscript"]):
                    tag.decompose()
                text = soup.get_text("\n", strip=True)
            return {"url": url, "text": text[:self.max_text_chars], "truncated": len(text) > self.max_text_chars, "citation_rule": "Cite this URL for facts taken from the fetched source."}
        raise EngineeringToolError("lookup mode must be search or fetch")

    def execute(self, name: str, arguments: dict[str, Any]) -> str:
        try:
            if name == "calculate_python":
                return _dumps(calculate_python(arguments))
            if name == "solve_math":
                return _dumps(solve_math(arguments))
            if name == "convert_units":
                return _dumps(convert_units(float(arguments["value"]), str(arguments["from_unit"]), str(arguments["to_unit"])))
            if name == "mechanical_calculator":
                return _dumps(mechanical_calculator(str(arguments["operation"]), dict(arguments.get("inputs") or {})))
            if name == "electrical_calculator":
                return _dumps(electrical_calculator(str(arguments["operation"]), dict(arguments.get("inputs") or {})))
            if name == "thermal_calculator":
                return _dumps(thermal_calculator(str(arguments["operation"]), dict(arguments.get("inputs") or {})))
            if name == "inspect_engineering_file":
                return _dumps(self.inspect_file(arguments))
            if name == "plot_engineering_data":
                return _dumps(self.plot_data(arguments))
            if name == "inspect_cad":
                return _dumps(self.inspect_cad(arguments))
            if name == "make_cad_model":
                return _dumps(self.make_cad_model(arguments))
            if name == "make_cad_part":
                return _dumps(self.make_cad_part(arguments))
            if name == "lookup_technical_reference":
                return _dumps(self.lookup_reference(arguments))
            raise EngineeringToolError(f"Unknown engineering tool: {name}")
        except Exception as exc:
            return f"Tool error: {type(exc).__name__}: {exc}"


ENGINEERING_TOOL_NAMES = {
    "calculate_python", "solve_math", "convert_units", "mechanical_calculator",
    "electrical_calculator", "thermal_calculator", "inspect_engineering_file",
    "plot_engineering_data", "inspect_cad", "make_cad_model", "make_cad_part", "lookup_technical_reference",
}


def _cad_catalog() -> str:
    from systems.cad_shapes import catalog
    return catalog()


def _cad_kinds() -> list[str]:
    from systems.cad_shapes import ALIASES, SHAPES
    return list(SHAPES) + list(ALIASES)


def engineering_tool_schemas() -> list[dict[str, Any]]:
    """Compact Ollama tool schemas; detailed operation keys live in descriptions."""
    return [
        {"type": "function", "function": {"name": "calculate_python", "description": "Safe deterministic numerical sandbox; no imports/files/network. operation: expression {expression,variables}; linear_solve {matrix,vector}; polyfit {x,y,degree}; interpolate {x,y,x_new}; integrate/root/minimize {expression,variable,x_min,x_max,variables}.", "parameters": {"type": "object", "required": ["operation"], "properties": {"operation": {"type": "string", "enum": ["expression","linear_solve","polyfit","interpolate","integrate","root","minimize"]}, "expression": {"type": "string"}, "variables": {"type": "object"}, "matrix": {"type": "array"}, "vector": {"type": "array"}, "x": {"type": "array"}, "y": {"type": "array"}, "x_new": {"type": "array"}, "degree": {"type": "integer"}, "variable": {"type": "string"}, "x_min": {"type": "number"}, "x_max": {"type": "number"}}}}},
        {"type": "function", "function": {"name": "solve_math", "description": "Symbolic algebra with SymPy. operation solve/simplify/differentiate/integrate/roots uses expression; system uses equations. Provide variable/symbols when useful.", "parameters": {"type": "object", "required": ["operation"], "properties": {"operation": {"type": "string", "enum": ["solve","simplify","differentiate","integrate","roots","system"]}, "expression": {"type": "string"}, "equations": {"type": "array", "items": {"type": "string"}}, "variable": {"type": "string"}, "symbols": {"type": "array", "items": {"type": "string"}}, "order": {"type": "integer"}}}}},
        {"type": "function", "function": {"name": "convert_units", "description": "Unit-aware converter for common engineering length/area/volume/mass/force/torque/pressure/energy/power/speed/time/electrical/frequency/temperature units.", "parameters": {"type": "object", "required": ["value","from_unit","to_unit"], "properties": {"value": {"type": "number"}, "from_unit": {"type": "string"}, "to_unit": {"type": "string"}}}}},
        {"type": "function", "function": {"name": "mechanical_calculator", "description": "Mechanical formulas. operation: gear_ratio/sprocket_ratio, vehicle_speed, rpm_from_speed, torque_multiplication, belt_chain_speed, pulley_cvt_ratio, beam, shaft_torsion, bearing_l10, bolt_preload_torque. Put named parameters in inputs; SI keys are explicit for beam/shaft/bearing/bolt.", "parameters": {"type": "object", "required": ["operation","inputs"], "properties": {"operation": {"type": "string", "enum": ["gear_ratio","sprocket_ratio","vehicle_speed","rpm_from_speed","torque_multiplication","belt_chain_speed","pulley_cvt_ratio","beam","shaft_torsion","bearing_l10","bolt_preload_torque"]}, "inputs": {"type": "object"}}}}},
        {"type": "function", "function": {"name": "electrical_calculator", "description": "Electrical formulas. operation: ohms_law, resistance_network, voltage_drop, wire_size_voltage_drop, battery_runtime, solar_array, dc_dc_sizing, rc_time_constant, capacitor_energy. Wire sizing is voltage-drop only and returns a code/ampacity warning.", "parameters": {"type": "object", "required": ["operation","inputs"], "properties": {"operation": {"type": "string", "enum": ["ohms_law","resistance_network","voltage_drop","wire_size_voltage_drop","battery_runtime","solar_array","dc_dc_sizing","rc_time_constant","capacitor_energy"]}, "inputs": {"type": "object"}}}}},
        {"type": "function", "function": {"name": "thermal_calculator", "description": "Thermal formulas. operation: temperature, heat_energy, thermal_resistance, steady_state_rise, conduction, convection. Use explicit SI-named input keys shown by the operation names.", "parameters": {"type": "object", "required": ["operation","inputs"], "properties": {"operation": {"type": "string", "enum": ["temperature","heat_energy","thermal_resistance","steady_state_rise","conduction","convection"]}, "inputs": {"type": "object"}}}}},
        {"type": "function", "function": {"name": "inspect_engineering_file", "description": "Read engineering data/document files without editing them. Supports txt/md/log/json/csv/tsv/xlsx/xlsm/pdf with bounded previews. Use exact local path from file discovery.", "parameters": {"type": "object", "required": ["path"], "properties": {"path": {"type": "string"}, "sheet": {"type": "string"}, "max_rows": {"type": "integer"}, "start_page": {"type": "integer"}, "page_count": {"type": "integer"}}}}},
        {"type": "function", "function": {"name": "plot_engineering_data", "description": "Create a PNG engineering plot from numeric series. series=[{x:[],y:[],label:''}]. Returns the saved path for later open/copy/export.", "parameters": {"type": "object", "required": ["series"], "properties": {"series": {"type": "array", "items": {"type": "object"}}, "title": {"type": "string"}, "x_label": {"type": "string"}, "y_label": {"type": "string"}, "grid": {"type": "boolean"}, "filename": {"type": "string"}, "dpi": {"type": "integer"}}}}},
        {"type": "function", "function": {"name": "inspect_cad", "description": "Deterministic CAD/mesh interrogation. mode=inspect gives bounds/area/volume; interference checks a second body; compare quantifies revision differences; section cuts a plane; overhang screens FDM support risk; draft screens face draft relative to a pull direction; thickness samples watertight wall thickness; enclosed_volume reports closed shells; mass_properties computes uniform-density mass/inertia. STEP advanced modes use local tessellation and report approximation warnings.", "parameters": {"type": "object", "required": ["path"], "properties": {"path": {"type": "string"}, "mode": {"type": "string", "enum": ["inspect","interference","compare","section","overhang","draft","thickness","enclosed_volume","mass_properties"]}, "other_path": {"type": "string"}, "plane_origin": {"type": "array", "items": {"type": "number"}}, "plane_normal": {"type": "array", "items": {"type": "number"}}, "build_direction": {"type": "array", "items": {"type": "number"}}, "max_overhang_from_vertical_deg": {"type": "number"}, "pull_direction": {"type": "array", "items": {"type": "number"}}, "minimum_draft_deg": {"type": "number"}, "sample_count": {"type": "integer"}, "density_kg_m3": {"type": "number"}, "native_length_unit": {"type": "string", "enum": ["mm","cm","m","in","ft"]}, "tessellation_tolerance": {"type": "number"}}}}},
        {"type": "function", "function": {"name": "make_cad_part", "description": "Make a simple 3D-printable part from one basic shape by giving its numbers (all mm); make_cad_model only for what these can't make. Shapes: " + _cad_catalog(), "parameters": {"type": "object", "required": ["shape", "name", "settings"], "properties": {"shape": {"type": "string", "enum": _cad_kinds()}, "name": {"type": "string", "description": "Short file name, e.g. cube_with_hole"}, "settings": {"type": "object", "description": "The shape's numbers by name, e.g. {\"width\": 20, \"depth\": 20, \"height\": 20, \"holes\": \"center\", \"hole_diameter\": 8}"}}}}},
        {"type": "function", "function": {"name": "make_cad_model", "description": "Make a 3D-printable STL from OpenSCAD source (1 unit = 1 mm; use $fn=64 for round parts). BMO saves name.scad and name.stl (never overwriting) and measures the STL. To change an earlier model, read its .scad with inspect_engineering_file and make it again.", "parameters": {"type": "object", "required": ["source", "name", "through_holes", "size_mm"], "properties": {"source": {"type": "string", "description": "Complete OpenSCAD design; no import/include/use/surface."}, "name": {"type": "string", "description": "Short file name without extension, e.g. cube_with_hole."}, "through_holes": {"type": "integer", "description": "How many holes go all the way through the part (0 if none); BMO checks the STL has exactly that many."}, "size_mm": {"type": "array", "items": {"type": "number"}, "description": "The part's outside size [width, depth, height] in mm, as asked; BMO checks it (any order, 0.5 mm leeway)."}}}}},
        {"type": "function", "function": {"name": "lookup_technical_reference", "description": "Engineering-only public technical-reference lookup. mode=search uses web search and returns URL/snippets; prefer manufacturer/standards domains. mode=fetch extracts bounded text from a public webpage/PDF. Cite returned URLs and separate published specs from calculations/inference.", "parameters": {"type": "object", "required": ["mode"], "properties": {"mode": {"type": "string", "enum": ["search","fetch"]}, "query": {"type": "string"}, "domains": {"type": "array", "items": {"type": "string"}}, "max_results": {"type": "integer"}, "url": {"type": "string"}}}}},
    ]
